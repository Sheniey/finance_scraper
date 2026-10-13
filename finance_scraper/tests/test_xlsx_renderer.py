from pathlib import Path

from openpyxl import load_workbook

from datatypes.ticker import Ticker
from finance.renderers.factory import RendererFactory
from finance.services.account.models import Account
from finance.services.ml.models import PredictionResult
from finance.services.reporter.models import Report
from finance.services.stock_analyzer.schema import StockAnalysis
from finance.types.report_mode import ReportMode


def _report() -> Report:
    return Report.model_construct(
        account=Account(name='Test User', email='test@example.com'),
        ticker=Ticker(ticker='AAPL'),
        news_compilation=[],
        analysis=StockAnalysis.model_construct(),
        prediction_result=PredictionResult.model_construct(
            ticker='=1+1',
        ),
        explanation=None,
        charts=True,
        mode=ReportMode.CLASSIC,
        custom_report=[],
    )


def test_xlsx_renderer_writes_report_sheets_and_escapes_formulas(
    tmp_path: Path,
) -> None:
    output_path: Path = tmp_path / 'reports' / 'AAPL.xlsx'
    renderer = RendererFactory.create('xlsx')

    renderer.render(_report(), str(output_path))

    workbook = load_workbook(output_path, read_only=True, data_only=False)
    assert workbook.sheetnames == ['Resumen', 'Datos']
    assert workbook['Resumen']['B2'].value == 'AAPL'
    formula_cell = next(
        row[1]
        for row in workbook['Datos'].iter_rows(min_row=2)
        if row[0].value == 'prediction_result.ticker'
    )
    assert formula_cell.value == "'=1+1"
    assert formula_cell.data_type == 's'
    workbook.close()


def test_xlsx_renderer_requires_a_filename() -> None:
    renderer = RendererFactory.create('xlsx')

    try:
        renderer.render(_report())
    except ValueError as error:
        assert str(error) == 'XLSX rendering requires an output filename'
    else:
        raise AssertionError('XLSX rendering should require an output filename')
