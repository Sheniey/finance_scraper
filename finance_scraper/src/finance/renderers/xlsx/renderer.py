import json
from collections.abc import Mapping
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from finance.renderers.utils import normalize_output
from finance.services.reporter import Report


_MAX_CELL_LENGTH: int = 32_767
_DANGEROUS_FORMULA_PREFIXES: tuple[str, ...] = ('=', '+', '-', '@', '\t', '\r')
_HEADER_FILL: PatternFill = PatternFill('solid', fgColor='0D1726')
_HEADER_FONT: Font = Font(color='FFFFFF', bold=True)
_TITLE_FONT: Font = Font(color='0D1726', bold=True, size=16)


def _flatten_value(
    value: Any,
    path: str = '',
) -> list[tuple[str, Any]]:
    if isinstance(value, Mapping):
        flattened: list[tuple[str, Any]] = []
        for key, nested_value in value.items():
            nested_path: str = f'{path}.{key}' if path else str(key)
            flattened.extend(_flatten_value(nested_value, nested_path))
        return flattened or [(path, '{}')]

    if isinstance(value, list):
        flattened = []
        for index, nested_value in enumerate(value):
            nested_path: str = f'{path}[{index}]'
            flattened.extend(_flatten_value(nested_value, nested_path))
        return flattened or [(path, '[]')]

    return [(path, value)]


def _cell_value(value: Any) -> str | int | float | bool | None:
    if value is None or isinstance(value, (bool, int, float)):
        return value

    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, default=str)

    text: str = str(value)
    if text.startswith(_DANGEROUS_FORMULA_PREFIXES):
        text = f"'{text}"
    if len(text) > _MAX_CELL_LENGTH:
        text = f'{text[:_MAX_CELL_LENGTH - 24]} [truncated]'
    return text


def _style_header(sheet: Any) -> None:
    for cell in sheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical='center')
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions


def _fit_columns(sheet: Any, maximum_width: int = 70) -> None:
    for column in sheet.columns:
        column_letter: str = get_column_letter(column[0].column)
        width: int = max(
            (
                len(str(cell.value))
                for cell in column
                if cell.value is not None
            ),
            default=10,
        )
        sheet.column_dimensions[column_letter].width = min(width + 2, maximum_width)


class XLSXRenderer:
    def render(self, report: Report, output: str | None = None) -> None:
        destination: str | None = normalize_output(output, report, 'xlsx')
        if destination is None:
            raise ValueError('XLSX rendering requires an output filename')

        payload: dict[str, Any] = report.model_dump(
            mode='json',
            exclude={'account'},
        )
        path: Path = Path(destination)
        path.parent.mkdir(parents=True, exist_ok=True)

        workbook: Workbook = Workbook()
        overview = workbook.active
        overview.title = 'Resumen'
        overview.append(['FINANCE SCRAPER'])
        overview['A1'].font = _TITLE_FONT
        overview.append(['Ticker', report.ticker.symbol.upper()])
        overview.append(['Modo', report.mode.value])
        overview.append([])
        overview.append(['Campo', 'Valor'])
        for cell in overview[5]:
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
        for field, value in (
            ('Noticias', len(report.news_compilation)),
            ('Gráficas incluidas', report.charts),
            ('Explicación IA', report.explanation is not None),
        ):
            overview.append([field, value])

        data_sheet = workbook.create_sheet('Datos')
        data_sheet.append(['Ruta', 'Valor'])
        for key, value in _flatten_value(payload):
            data_sheet.append([_cell_value(key), _cell_value(value)])
        _style_header(data_sheet)
        _fit_columns(data_sheet)
        _fit_columns(overview, maximum_width=50)

        workbook.save(path)

    def normalize_output(
        self,
        report: Report,
        output: str | None = None,
    ) -> Path:
        destination: str | None = normalize_output(output, report, 'xlsx')
        if destination is None:
            raise ValueError('XLSX rendering requires an output filename')
        return Path(destination)
